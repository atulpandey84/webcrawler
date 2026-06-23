from pptx import Presentation
from pptx.util import Inches, Pt

def create_presentation():
    prs = Presentation()

    # Slide 1: Title
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "Enterprise Infrastructure Migration to GCP: Solution Architecture"
    subtitle.text = "End-to-End Solution Architecture and Landing Zone Design\nPrepared by Jules, Senior Software Engineer"

    # Slide 2: Executive Summary
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "Executive Summary"
    tf = slide.shapes.placeholders[1].text_frame
    tf.text = "Objective: Transition enterprise workloads from on-prem to Google Cloud Platform."
    p = tf.add_paragraph()
    p.text = "Core Strategy: Establish a robust, secure, and scalable Landing Zone."
    p = tf.add_paragraph()
    p.text = "Key Pillars:"
    for item in ["Centralized Governance (Organization & Folders)", "Networking (Shared VPC & Hybrid Connectivity)", "Security (IAM, Encryption, Edge Security)", "Operations (Cloud Logging & Monitoring)"]:
        sp = tf.add_paragraph()
        sp.text = item
        sp.level = 1

    # Slide 3: GCP Organization Hierarchy
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "GCP Organization Hierarchy"
    tf = slide.shapes.placeholders[1].text_frame
    tf.text = "Organization Node: Root of the hierarchy."
    p = tf.add_paragraph()
    p.text = "Folders: Common, Production, Non-Production, Sandbox."
    p = tf.add_paragraph()
    p.text = "Projects:"
    for item in ["Host Project: Manages Shared VPC", "Service Projects: Application-specific resources", "Logging/Audit Project: Centralized logs", "Security Project: Secrets management, KMS"]:
        sp = tf.add_paragraph()
        sp.text = item
        sp.level = 1

    # Slide 4: GCP Landing Zone & Automation
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "GCP Landing Zone & Automation"
    tf = slide.shapes.placeholders[1].text_frame
    tf.text = "Infrastructure as Code (IaC): Terraform for repeatable deployments."
    p = tf.add_paragraph()
    p.text = "Landing Zone Components:"
    for item in ["Resource Hierarchy setup", "Identity & Access Management", "Networking (Hub-and-Spoke or Shared VPC)", "Logging & Monitoring sinks", "Security Guardrails (Organization Policies)"]:
        sp = tf.add_paragraph()
        sp.text = item
        sp.level = 1

    # Slide 5: Networking: Shared VPC & Hybrid Connectivity
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "Networking: Shared VPC & Hybrid Connectivity"
    tf = slide.shapes.placeholders[1].text_frame
    tf.text = "Shared VPC: Centralized management of network resources in a Host Project."
    p = tf.add_paragraph()
    p.text = "Service Projects: Consume subnets from the Host Project."
    p = tf.add_paragraph()
    p.text = "Hybrid Connectivity: Cloud Interconnect (High Bandwidth), Cloud VPN (Encrypted)."
    p = tf.add_paragraph()
    p.text = "Additional Services: Cloud DNS (Internal resolution), Cloud NAT (Private VM internet)."

    # Slide 6: IAM & Security Architecture
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "IAM & Security Architecture"
    tf = slide.shapes.placeholders[1].text_frame
    tf.text = "Identity Management: Syncing On-Prem AD with Cloud Identity."
    p = tf.add_paragraph()
    p.text = "IAM Principles: Least Privilege, Custom Roles, Service Account security."
    p = tf.add_paragraph()
    p.text = "Network Security: Cloud Armor (WAF), IAP (Remote Access), VPC Service Controls."
    p = tf.add_paragraph()
    p.text = "Data Security: Cloud KMS, CMEK (Customer-Managed Encryption Keys)."

    # Slide 7: Centralized Logging & Monitoring
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "Centralized Logging & Monitoring"
    tf = slide.shapes.placeholders[1].text_frame
    tf.text = "Cloud Logging:"
    for item in ["Log Sinks: Exporting to BigQuery or GCS", "Centralized Logging Project for audit trails"]:
        sp = tf.add_paragraph()
        sp.text = item
        sp.level = 1
    p = tf.add_paragraph()
    p.text = "Cloud Monitoring:"
    for item in ["Multi-project dashboards", "Uptime checks and Alerts (Slack/Email)"]:
        sp = tf.add_paragraph()
        sp.text = item
        sp.level = 1

    # Slide 8: Governance & Compliance
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "Governance & Compliance"
    tf = slide.shapes.placeholders[1].text_frame
    tf.text = "Organization Policy Service: Restrict locations, disable external IPs."
    p = tf.add_paragraph()
    p.text = "Policy Intelligence: Recommenders for IAM and firewall rules."
    p = tf.add_paragraph()
    p.text = "Compliance: Mapping to SOC2, HIPAA, or GDPR."
    p = tf.add_paragraph()
    p.text = "Resource Quotas: Managing limits across projects."

    # Slide 9: Migration Strategy & Phases
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "Migration Strategy & Phases"
    tf = slide.shapes.placeholders[1].text_frame
    tf.text = "Phase 1: Assess (Discovery, TCO)."
    p = tf.add_paragraph()
    p.text = "Phase 2: Plan (Architecture design, Landing Zone setup)."
    p = tf.add_paragraph()
    p.text = "Phase 3: Deploy (Pilot migration, Wave-based migration)."
    p = tf.add_paragraph()
    p.text = "Phase 4: Optimize (Right-sizing, Modernization)."
    p = tf.add_paragraph()
    p.text = "Tools: Migrate for Compute Engine (M4CE), Anthos."

    # Slide 10: Conclusion
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "Conclusion"
    tf = slide.shapes.placeholders[1].text_frame
    tf.text = "Summary: A robust GCP Landing Zone provides a scalable, secure, and governed environment for enterprise workloads."
    p = tf.add_paragraph()
    p.text = "Next Steps: Begin Assessment phase and POC for Landing Zone components."

    prs.save("GCP_Migration_Architecture.pptx")
    print("Presentation saved as GCP_Migration_Architecture.pptx")

if __name__ == "__main__":
    create_presentation()
